"""Convert markdown-like resume text into a formatted .docx file."""

from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.shared import Pt


def build_docx(content: str, title: str = "Resume") -> bytes:
    """Parse *content* (simple markdown) and return .docx bytes."""
    doc = Document()

    style = doc.styles["Normal"]
    style.font.size = Pt(11)
    style.font.name = "Calibri"

    for line in content.splitlines():
        stripped = line.strip()

        if not stripped:
            # Empty line → paragraph break
            doc.add_paragraph("")
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        elif stripped.startswith(("- ", "* ")):
            _add_rich_paragraph(doc.add_paragraph(style="List Bullet"), stripped[2:])
        else:
            _add_rich_paragraph(doc.add_paragraph(), stripped)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_rich_paragraph(paragraph, text: str) -> None:
    """Add runs to *paragraph*, converting **bold** markers."""
    parts = _BOLD_RE.split(text)
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = paragraph.add_run(part)
        if idx % 2 == 1:
            run.bold = True
