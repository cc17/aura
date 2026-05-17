"""Tests for the docx builder utility."""

from io import BytesIO

from docx import Document

from backend.utils.docx_builder import build_docx


def _load(docx_bytes: bytes) -> Document:
    return Document(BytesIO(docx_bytes))


def test_basic_text():
    result = build_docx("Hello world")
    doc = _load(result)
    assert any("Hello world" in p.text for p in doc.paragraphs)


def test_headings():
    content = "# Title\n## Section\n### Subsection\nBody text"
    doc = _load(build_docx(content))

    texts = [(p.text, p.style.name) for p in doc.paragraphs]
    assert ("Title", "Heading 1") in texts
    assert ("Section", "Heading 2") in texts
    assert ("Subsection", "Heading 3") in texts
    assert any("Body text" in p.text for p in doc.paragraphs)


def test_bold_text():
    doc = _load(build_docx("This is **bold** text"))
    para = next(p for p in doc.paragraphs if "bold" in p.text)
    bold_runs = [r for r in para.runs if r.bold]
    assert any("bold" in r.text for r in bold_runs)


def test_list_items():
    content = "- Item one\n* Item two"
    doc = _load(build_docx(content))
    list_paras = [p for p in doc.paragraphs if "List" in p.style.name]
    assert len(list_paras) == 2


def test_empty_content():
    result = build_docx("")
    doc = _load(result)
    # Should produce a valid doc with no meaningful text
    assert doc is not None


def test_returns_bytes():
    result = build_docx("test")
    assert isinstance(result, bytes)
    assert len(result) > 0
